"""Decode content into native sampled fields, never route by filename suffix."""

from fractions import Fraction
from zipfile import BadZipFile, ZipFile

import av
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .calculus import difference, numeric


class Unmodeled(Exception):
    """Keep the byte field and explain why no decoded field was published."""


def rational(value):
    value = Fraction(value)
    return [value.numerator, value.denominator]


class ContentReader:
    """Do not expose a filename/extension to FFmpeg's format probing."""

    def __init__(self, stream):
        self.stream = stream

    def read(self, size=-1):
        return self.stream.read(size)

    def seek(self, offset, whence=0):
        return self.stream.seek(offset, whence)

    def tell(self):
        return self.stream.tell()


def video_planes(frame):
    """Strip allocator padding; preserve native color planes and bit depth."""
    fmt = frame.format
    packed = {"rgb24": ("u1", 3), "bgr24": ("u1", 3),
              "rgba": ("u1", 4), "bgra": ("u1", 4),
              "argb": ("u1", 4), "abgr": ("u1", 4),
              "rgb48be": (">u2", 3), "rgb48le": ("<u2", 3),
              "rgba64be": (">u2", 4), "rgba64le": ("<u2", 4)}
    if fmt.name in packed:
        dtype, channels = packed[fmt.name]
        layouts = [(np.dtype(dtype), channels)]
    elif (fmt.is_planar or fmt.name.startswith("gray")) and not fmt.is_bit_stream:
        layouts = []
        for i in range(len(frame.planes)):
            components = [c for c in fmt.components if c.plane == i]
            if len(components) != 1 or components[0].bits not in (8, 9, 10, 12, 14, 16):
                raise Unmodeled("unsupported_pixel_layout:" + fmt.name)
            bits = components[0].bits
            layouts.append((np.dtype("u1" if bits == 8 else ">u2" if fmt.is_big_endian else "<u2"), 1))
    else:
        raise Unmodeled("unsupported_pixel_layout:" + fmt.name)
    arrays = []
    for plane, (dtype, channels) in zip(frame.planes, layouts):
        if plane.line_size <= 0 or plane.line_size % dtype.itemsize:
            raise Unmodeled("unsupported_plane_stride")
        rows = np.frombuffer(plane, dtype=dtype, count=plane.height * plane.line_size // dtype.itemsize)
        rows = rows.reshape(plane.height, plane.line_size // dtype.itemsize)
        values = rows[:, :plane.width * channels].copy()
        arrays.append(values.reshape(plane.height, plane.width, channels) if channels > 1
                      else values.reshape(plane.height, plane.width))
    return arrays


def decode_pdf(incoming, writer):
    """Text code points per page; no OCR or inferred coordinates."""
    try:
        reader = PdfReader(incoming)
        if reader.is_encrypted:
            raise Unmodeled("encrypted_pdf")
        if len(reader.pages) > writer.max_frames:
            raise Unmodeled("page_limit")
        fields = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if len(text) > writer.max_elements:
                raise Unmodeled("array_limit")
            values = np.fromiter((ord(char) for char in text), dtype=np.int32, count=len(text))
            fields.append({"id": f"page_{index + 1}_text", "axes": ["codepoint"],
                           "value_unit": "unicode_codepoint", "page": index + 1,
                           "time": {"kind": "static"}, "chunks": [
                               {"samples": writer.array(values),
                                "temporal": {"kind": "constant", "rate": 0}}]})
        if not any(field["chunks"][0]["samples"]["shape"][0] for field in fields):
            raise Unmodeled("no_extractable_pdf_text")
        return {"adapter": "pypdf", "status": "decoded", "unmodeled_streams": [],
                "fields": fields}
    except (PdfReadError, OSError, KeyError, TypeError) as exc:
        raise Unmodeled("invalid_pdf") from exc


def decode_xlsx(incoming, writer):
    """Numeric cell grids; reject formulas and other cell types explicitly."""
    try:
        with ZipFile(incoming) as archive:
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise Unmodeled("unsupported_zip_container")
            if (len(names) > 1000 or
                    sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024):
                raise Unmodeled("xlsx_uncompressed_limit")
        incoming.seek(0)
        book = load_workbook(incoming, read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) > writer.max_frames:
                raise Unmodeled("sheet_limit")
            fields = []
            for index, sheet in enumerate(book.worksheets):
                rows, cols = sheet.max_row or 0, sheet.max_column or 0
                if rows * cols > writer.max_elements:
                    raise Unmodeled("array_limit")
                if not rows or not cols:
                    continue
                values = np.empty((rows, cols), dtype=np.float64)
                for row in sheet.iter_rows():
                    for cell in row:
                        value = cell.value
                        if (cell.data_type == "f" or isinstance(value, bool) or
                                not isinstance(value, (int, float)) or
                                not np.isfinite(value)):
                            raise Unmodeled("non_numeric_xlsx_cell")
                        if isinstance(value, int) and abs(value) > 2**53:
                            raise Unmodeled("xlsx_integer_precision_limit")
                        values[cell.row - 1, cell.column - 1] = value
                fields.append({"id": f"sheet_{index + 1}", "axes": ["row", "column"],
                               "value_unit": "cell_number", "sheet": sheet.title,
                               "time": {"kind": "static"}, "chunks": [
                                   {"samples": writer.array(values),
                                    "temporal": {"kind": "constant", "rate": 0}}]})
            if not fields:
                raise Unmodeled("empty_xlsx")
            return {"adapter": "openpyxl", "status": "decoded",
                    "unmodeled_streams": [], "fields": fields}
        finally:
            book.close()
    except (BadZipFile, InvalidFileException, OSError, KeyError, ValueError, TypeError) as exc:
        raise Unmodeled("invalid_xlsx") from exc


def decode(source, writer):
    with source.open("rb") as incoming:
        header = incoming.read(6)
        if header.startswith(b"%PDF-"):
            incoming.seek(0)
            return decode_pdf(incoming, writer)
        if header.startswith(b"PK"):
            incoming.seek(0)
            return decode_xlsx(incoming, writer)
        if header == b"\x93NUMPY":
            array = np.load(source, mmap_mode="r", allow_pickle=False)
            if array.size > writer.max_elements:
                raise Unmodeled("array_limit")
            numeric(array)
            return {"adapter": "numpy", "status": "decoded", "unmodeled_streams": [],
                    "fields": [{"id": "tensor", "axes": [f"index_{i}" for i in range(array.ndim)],
                                "time": {"kind": "unmodeled"},
                                "chunks": [{"samples": writer.array(array)}]}]}
        incoming.seek(0)
        # Custom I/O only. A playlist cannot fetch network or local dependencies.
        with av.open(ContentReader(incoming), options={"protocol_whitelist": "pipe"}) as container:
            if any(x in container.format.name.split(",") for x in ("hls", "dash", "concat", "image2")):
                raise Unmodeled("external_reference_container")
            streams = [s for s in container.streams if s.type in ("audio", "video")]
            if not streams:
                raise Unmodeled("no_supported_media_stream")
            for stream in streams:
                stream.codec_context.options = {**stream.codec_context.options, "err_detect": "explode"}
            omitted = [{"index": s.index, "type": s.type} for s in container.streams if s not in streams]
            static = container.format.name in ("png_pipe", "jpeg_pipe", "bmp_pipe", "tiff_pipe", "webp_pipe")
            fields, previous, signatures, counts = {}, {}, {}, {}
            for packet in container.demux(streams):
                for frame in packet.decode():
                    stream = packet.stream
                    if getattr(frame, "is_corrupt", False):
                        raise Unmodeled("corrupt_decoded_frame")
                    counts[stream.index] = counts.get(stream.index, 0) + 1
                    if counts[stream.index] > writer.max_frames:
                        raise Unmodeled("frame_limit")
                    is_static = static and stream.type == "video"
                    if is_static and counts[stream.index] > 1:
                        raise Unmodeled("multiple_frames_without_declared_clock")
                    timestamp = (Fraction(frame.pts) * frame.time_base
                                 if frame.pts is not None and frame.time_base is not None else None)
                    if not is_static and timestamp is None:
                        raise Unmodeled("missing_timestamp")
                    if stream.type == "video":
                        if frame.width * frame.height * 4 > writer.max_elements:
                            raise Unmodeled("array_limit")
                        arrays = video_planes(frame)
                        signature = (frame.format.name, frame.width, frame.height)
                        metadata = {"pixel_format": frame.format.name,
                                    "component_bits": [c.bits for c in frame.format.components],
                                    "colorspace": int(frame.colorspace), "color_range": int(frame.color_range),
                                    "stream_color_primaries": int(stream.codec_context.color_primaries),
                                    "stream_color_transfer": int(stream.codec_context.color_trc),
                                    "stream_sample_aspect_ratio": (rational(stream.sample_aspect_ratio)
                                                                   if stream.sample_aspect_ratio else None),
                                    "width": frame.width, "height": frame.height}
                        step = None
                    else:
                        channels = len(frame.layout.channels)
                        if frame.samples * channels > writer.max_elements:
                            raise Unmodeled("array_limit")
                        values = frame.to_ndarray()
                        values = values.T.copy() if frame.format.is_planar else values.reshape(-1, channels).copy()
                        arrays = [values]
                        signature = (frame.format.name, frame.layout.name, frame.sample_rate)
                        metadata = {"sample_format": frame.format.name, "sample_rate": frame.sample_rate,
                                    "channels": [c.name for c in frame.layout.channels]}
                        step = Fraction(1, frame.sample_rate)
                    if stream.index in signatures and signature != signatures[stream.index]:
                        raise Unmodeled("stream_configuration_changed")
                    signatures[stream.index] = signature
                    for plane_index, values in enumerate(arrays):
                        numeric(values)
                        key = f"stream_{stream.index}_plane_{plane_index}"
                        axes = (["sample", "channel"] if step else
                                ["y", "x"] + (["component"] if values.ndim == 3 else []))
                        field = fields.setdefault(key, {"id": key, "axes": axes,
                            "value_unit": "native_sample", "spatial_unit": "plane_pixel" if not step else None,
                            "native": metadata, "time": {"kind": "static" if is_static else "sampled",
                            "unit": "second", "interpolation": "none"}, "chunks": []})
                        chunk = {"samples": writer.array(values)}
                        if is_static:
                            chunk["temporal"] = {"kind": "constant", "rate": 0}
                        else:
                            chunk["start"] = rational(timestamp)
                            chunk["native_pts"] = frame.pts
                            chunk["native_time_base"] = rational(frame.time_base)
                            duration = getattr(frame, "duration", 0)
                            chunk["duration"] = (rational(len(values) * step) if step else
                                                 rational(duration * frame.time_base) if duration else None)
                            if step:
                                chunk["step"] = rational(step)
                                chunk["temporal"] = {"kind": "forward_difference", "dt": rational(step),
                                    "delta": writer.array(difference(values[:-1], values[1:]))}
                            if key in previous:
                                last_values, last_time = previous[key]
                                dt = timestamp - last_time
                                if dt <= 0:
                                    raise Unmodeled("nonincreasing_timestamps")
                                delta = difference(last_values, values[0] if step else values)
                                chunk["from_previous"] = {"kind": "finite_difference", "dt": rational(dt),
                                                          "delta": writer.array(delta)}
                            previous[key] = ((values[-1].copy(), timestamp + (len(values) - 1) * step)
                                             if step else (values, timestamp))
                        field["chunks"].append(chunk)
            if not fields or any(s.index not in counts for s in streams):
                raise Unmodeled("empty_or_undecodable_stream")
            return {"adapter": "pyav", "decoder_version": av.__version__,
                    "container": container.format.name,
                    "status": "partial" if omitted else "decoded", "unmodeled_streams": omitted,
                    "fields": list(fields.values())}
